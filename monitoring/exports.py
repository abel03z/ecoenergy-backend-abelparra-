"""Exportación de mantenciones a Excel (.xlsx) con openpyxl."""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADERS = [
    "ID", "Dispositivo", "Zona", "Departamento", "Organización",
    "Tipo", "Estado", "Técnico", "Programada", "Finalizada", "Creada",
]


def safe_text(value):
    """
    Evita inyección de fórmulas: un texto que empieza con = + - @ sería
    interpretado por Excel como fórmula. Se le antepone una comilla simple.
    """
    text = "" if value is None else str(value)
    if text and text[0] in ("=", "+", "-", "@"):
        return "'" + text
    return text


def excel_datetime(value):
    """openpyxl no admite datetimes con zona horaria: se pasa a hora local sin tzinfo."""
    if value is None:
        return None
    return timezone.localtime(value).replace(tzinfo=None, microsecond=0)


def build_maintenance_workbook(queryset):
    """
    Genera el libro a partir de un queryset YA filtrado por permisos, scoping
    y borrado lógico (lo arma la vista). Los datos salen de la base de datos.
    """
    queryset = queryset.select_related(
        "device__zone__department__organization", "technician"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Mantenciones"
    ws.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for m in queryset.iterator(chunk_size=500):
        zone = m.device.zone
        ws.append([
            m.pk,
            safe_text(m.device.name),
            safe_text(zone.name),
            safe_text(zone.department.name),
            safe_text(zone.department.organization.name),
            m.get_type_display(),
            m.get_status_display(),
            safe_text(m.technician.get_username()) if m.technician else "",
            excel_datetime(m.scheduled_at),
            excel_datetime(m.completed_at),
            excel_datetime(m.created_at),
        ])

    date_format = "dd/mm/yyyy hh:mm"
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=11):
        for cell in row:
            cell.number_format = date_format

    widths = [8, 32, 28, 22, 20, 14, 14, 16, 18, 18, 18]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
