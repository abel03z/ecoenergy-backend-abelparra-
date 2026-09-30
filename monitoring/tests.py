from datetime import timedelta

from django.contrib.auth.models import Group, Permission, User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import UserProfile
from devices.models import Category, Device, Manufacturer
from organizations.models import Department, Organization, Zone

from .models import Maintenance


def make_user(username, org, group):
    user = User.objects.create_user(username, password="x")
    user.groups.add(group)
    UserProfile.objects.create(user=user, organization=org, employee_code=f"E-{username}")
    return user


def make_group(name, actions):
    group = Group.objects.create(name=name)
    for action in actions:
        group.permissions.add(Permission.objects.get(codename=f"{action}_maintenance"))
    return group


def fmt(dt):
    return dt.strftime("%Y-%m-%dT%H:%M")


class MaintenanceCrudTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.norte = Organization.objects.create(name="Norte")
        cls.sur = Organization.objects.create(name="Sur")
        d_norte = Department.objects.create(organization=cls.norte, name="Ops N")
        d_sur = Department.objects.create(organization=cls.sur, name="Ops S")
        zone_norte = Zone.objects.create(department=d_norte, name="ZN", consumption_limit=10)
        zone_sur = Zone.objects.create(department=d_sur, name="ZS", consumption_limit=10)
        category = Category.objects.create(name="Medidor")
        maker = Manufacturer.objects.create(name="Acme")
        cls.dev_norte = Device.objects.create(name="Dev Norte", zone=zone_norte, category=category, manufacturer=maker)
        cls.dev_sur = Device.objects.create(name="Dev Sur", zone=zone_sur, category=category, manufacturer=maker)

        full = make_group("Admin mantenciones", ("add", "change", "delete", "view"))
        operator = make_group("Operador mantenciones", ("add", "change", "view"))
        reader = make_group("Lector mantenciones", ("view",))
        cls.admin = make_user("m_admin", cls.norte, full)
        cls.operator = make_user("m_operator", cls.norte, operator)
        cls.reader = make_user("m_reader", cls.norte, reader)
        cls.tech_sur = make_user("tech_sur", cls.sur, reader)

        cls.when = timezone.now().replace(second=0, microsecond=0) + timedelta(days=3)
        cls.m_norte = Maintenance.objects.create(
            device=cls.dev_norte, type=Maintenance.Type.PREVENTIVE, scheduled_at=cls.when
        )
        cls.m_sur = Maintenance.objects.create(
            device=cls.dev_sur, type=Maintenance.Type.PREVENTIVE, scheduled_at=cls.when
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def _payload(self, **extra):
        data = {
            "device": self.dev_norte.pk,
            "technician": "",
            "type": Maintenance.Type.CORRECTIVE,
            "status": Maintenance.Status.PENDING,
            "scheduled_at": fmt(timezone.now() + timedelta(days=5)),
            "completed_at": "",
        }
        data.update(extra)
        return data

    # --- acceso y permisos diferenciados
    def test_requiere_login(self):
        self.client.logout()
        r = self.client.get(reverse("monitoring:maintenance_list"))
        self.assertEqual(r.status_code, 302)

    def test_lector_solo_puede_ver(self):
        self.client.force_login(self.reader)
        self.assertEqual(self.client.get(reverse("monitoring:maintenance_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("monitoring:maintenance_create")).status_code, 403)
        self.assertEqual(self.client.post(reverse("monitoring:maintenance_create"), self._payload()).status_code, 403)
        self.assertEqual(self.client.get(reverse("monitoring:maintenance_update", args=[self.m_norte.pk])).status_code, 403)
        r = self.client.post(reverse("monitoring:maintenance_delete", args=[self.m_norte.pk]))
        self.assertEqual(r.status_code, 403)

    def test_operador_crea_y_edita_pero_no_elimina(self):
        self.client.force_login(self.operator)
        r = self.client.post(reverse("monitoring:maintenance_create"), self._payload())
        self.assertRedirects(r, reverse("monitoring:maintenance_list"))
        r = self.client.post(reverse("monitoring:maintenance_delete", args=[self.m_norte.pk]))
        self.assertEqual(r.status_code, 403)
        self.assertTrue(Maintenance.objects.filter(pk=self.m_norte.pk).exists())

    # --- create / update
    def test_crear_mantencion(self):
        r = self.client.post(reverse("monitoring:maintenance_create"), self._payload())
        self.assertRedirects(r, reverse("monitoring:maintenance_list"))
        self.assertEqual(Maintenance.objects.filter(device=self.dev_norte).count(), 2)

    def test_actualizar_mantencion(self):
        r = self.client.post(
            reverse("monitoring:maintenance_update", args=[self.m_norte.pk]),
            self._payload(type=Maintenance.Type.PREVENTIVE, scheduled_at=fmt(self.when),
                          status=Maintenance.Status.IN_PROGRESS),
        )
        self.assertRedirects(r, reverse("monitoring:maintenance_list"))
        self.m_norte.refresh_from_db()
        self.assertEqual(self.m_norte.status, Maintenance.Status.IN_PROGRESS)

    # --- validaciones
    def test_pendiente_nueva_no_puede_ir_en_el_pasado(self):
        r = self.client.post(
            reverse("monitoring:maintenance_create"),
            self._payload(scheduled_at=fmt(timezone.now() - timedelta(days=1))),
        )
        self.assertEqual(r.status_code, 200)
        self.assertIn("scheduled_at", r.context["form"].errors)

    def test_estado_done_exige_fecha_de_termino(self):
        r = self.client.post(
            reverse("monitoring:maintenance_create"),
            self._payload(status=Maintenance.Status.DONE),
        )
        self.assertIn("completed_at", r.context["form"].errors)

    def test_fecha_de_termino_solo_si_esta_done(self):
        r = self.client.post(
            reverse("monitoring:maintenance_create"),
            self._payload(completed_at=fmt(timezone.now() - timedelta(hours=1))),
        )
        self.assertIn("completed_at", r.context["form"].errors)

    def test_fecha_de_termino_no_puede_ser_futura(self):
        r = self.client.post(
            reverse("monitoring:maintenance_create"),
            self._payload(status=Maintenance.Status.DONE,
                          completed_at=fmt(timezone.now() + timedelta(days=2))),
        )
        self.assertIn("completed_at", r.context["form"].errors)

    def test_done_con_fecha_valida_se_guarda(self):
        past = timezone.now() - timedelta(days=2)
        r = self.client.post(
            reverse("monitoring:maintenance_create"),
            self._payload(status=Maintenance.Status.DONE, scheduled_at=fmt(past),
                          completed_at=fmt(timezone.now() - timedelta(hours=1))),
        )
        self.assertRedirects(r, reverse("monitoring:maintenance_list"))

    def test_mantencion_duplicada(self):
        r = self.client.post(
            reverse("monitoring:maintenance_create"),
            self._payload(type=Maintenance.Type.PREVENTIVE, scheduled_at=fmt(self.when)),
        )
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.context["form"].non_field_errors())

    # --- scoping
    def test_listado_solo_muestra_mantenciones_de_mi_organizacion(self):
        r = self.client.get(reverse("monitoring:maintenance_list"))
        self.assertContains(r, "Dev Norte")
        self.assertNotContains(r, "Dev Sur")

    def test_no_puede_editar_ni_eliminar_de_otra_organizacion(self):
        self.assertEqual(self.client.get(reverse("monitoring:maintenance_update", args=[self.m_sur.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse("monitoring:maintenance_delete", args=[self.m_sur.pk])).status_code, 404)
        self.assertTrue(Maintenance.objects.filter(pk=self.m_sur.pk).exists())

    def test_no_puede_usar_dispositivo_ni_tecnico_de_otra_organizacion(self):
        r = self.client.post(reverse("monitoring:maintenance_create"), self._payload(device=self.dev_sur.pk))
        self.assertIn("device", r.context["form"].errors)
        r = self.client.post(reverse("monitoring:maintenance_create"), self._payload(technician=self.tech_sur.pk))
        self.assertIn("technician", r.context["form"].errors)

    # --- borrado lógico
    def test_eliminar_es_borrado_logico(self):
        r = self.client.post(reverse("monitoring:maintenance_delete", args=[self.m_norte.pk]))
        self.assertRedirects(r, reverse("monitoring:maintenance_list"))
        self.assertFalse(Maintenance.objects.filter(pk=self.m_norte.pk).exists())
        self.assertIsNotNone(Maintenance.all_objects.get(pk=self.m_norte.pk).deleted_at)
        self.assertNotContains(self.client.get(reverse("monitoring:maintenance_list")), "Dev Norte")

    def test_eliminar_por_get_no_esta_permitido(self):
        r = self.client.get(reverse("monitoring:maintenance_delete", args=[self.m_norte.pk]))
        self.assertEqual(r.status_code, 405)
        self.assertTrue(Maintenance.objects.filter(pk=self.m_norte.pk).exists())


class MaintenanceExcelExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.norte = Organization.objects.create(name="Norte")
        cls.sur = Organization.objects.create(name="Sur")
        d_norte = Department.objects.create(organization=cls.norte, name="Ops N")
        d_sur = Department.objects.create(organization=cls.sur, name="Ops S")
        zone_norte = Zone.objects.create(department=d_norte, name="ZN", consumption_limit=10)
        zone_sur = Zone.objects.create(department=d_sur, name="ZS", consumption_limit=10)
        category = Category.objects.create(name="Medidor")
        maker = Manufacturer.objects.create(name="Acme")
        cls.dev_norte = Device.objects.create(name="Dev Norte", zone=zone_norte, category=category, manufacturer=maker)
        cls.dev_sur = Device.objects.create(name="Dev Sur", zone=zone_sur, category=category, manufacturer=maker)
        cls.evil = Device.objects.create(name="=HYPERLINK(\"http://x\")", zone=zone_norte, category=category, manufacturer=maker)

        viewer = Group.objects.create(name="Con vista")
        viewer.permissions.add(Permission.objects.get(codename="view_maintenance"))
        cls.viewer = make_user("x_viewer", cls.norte, viewer)
        cls.no_perm = make_user("x_noperm", cls.norte, Group.objects.create(name="Sin permisos"))
        cls.viewer_sur = make_user("x_viewer_sur", cls.sur, viewer)
        cls.superuser = User.objects.create_superuser("x_root", password="x")

        when = timezone.now() + timedelta(days=2)
        for i in range(3):
            Maintenance.objects.create(device=cls.dev_norte, type=Maintenance.Type.PREVENTIVE,
                                       scheduled_at=when + timedelta(hours=i))
        cls.deleted = Maintenance.objects.create(device=cls.dev_norte, type=Maintenance.Type.CORRECTIVE,
                                                 scheduled_at=when + timedelta(days=9))
        cls.deleted.soft_delete()
        Maintenance.objects.create(device=cls.dev_sur, type=Maintenance.Type.CORRECTIVE, scheduled_at=when)
        Maintenance.objects.create(device=cls.evil, type=Maintenance.Type.PREVENTIVE, scheduled_at=when)

    def _workbook(self, response):
        from io import BytesIO
        from openpyxl import load_workbook
        return load_workbook(BytesIO(response.content)).active

    def _rows(self, response):
        ws = self._workbook(response)
        return [[c.value for c in row] for row in ws.iter_rows(min_row=2)]

    def test_requiere_login(self):
        r = self.client.get(reverse("monitoring:maintenance_export"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r.url)

    def test_requiere_permiso(self):
        self.client.force_login(self.no_perm)
        r = self.client.get(reverse("monitoring:maintenance_export"))
        self.assertEqual(r.status_code, 403)

    def test_descarga_un_xlsx_real(self):
        self.client.force_login(self.viewer)
        r = self.client.get(reverse("monitoring:maintenance_export"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn(".xlsx", r["Content-Disposition"])
        self.assertEqual(r.content[:2], b"PK")  # un .xlsx es un zip

    def test_tiene_encabezados_y_datos_desde_la_base(self):
        self.client.force_login(self.viewer)
        r = self.client.get(reverse("monitoring:maintenance_export"))
        ws = self._workbook(r)
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[:3], ["ID", "Dispositivo", "Zona"])
        self.assertIn("Estado", headers)
        rows = self._rows(r)
        self.assertGreater(len(rows), 0)
        ids = {row[0] for row in rows}
        self.assertTrue(ids <= set(Maintenance.objects.values_list("pk", flat=True)))

    def test_respeta_scoping_y_borrado_logico(self):
        self.client.force_login(self.viewer)
        rows = self._rows(self.client.get(reverse("monitoring:maintenance_export")))
        ids = {row[0] for row in rows}
        devices = {row[1] for row in rows}
        # 3 normales + 1 del dispositivo "evil" (ambos de Norte); sin la eliminada ni las de Sur
        self.assertEqual(len(rows), 4)
        self.assertNotIn(self.deleted.pk, ids)
        self.assertNotIn("Dev Sur", devices)
        self.assertEqual({row[4] for row in rows}, {"Norte"})

    def test_otra_organizacion_solo_ve_las_suyas(self):
        self.client.force_login(self.viewer_sur)
        rows = self._rows(self.client.get(reverse("monitoring:maintenance_export")))
        self.assertEqual([row[1] for row in rows], ["Dev Sur"])

    def test_superusuario_exporta_todo_lo_vigente(self):
        self.client.force_login(self.superuser)
        rows = self._rows(self.client.get(reverse("monitoring:maintenance_export")))
        self.assertEqual(len(rows), Maintenance.objects.count())
        self.assertNotIn(self.deleted.pk, {row[0] for row in rows})

    def test_neutraliza_formulas_en_el_texto(self):
        self.client.force_login(self.viewer)
        rows = self._rows(self.client.get(reverse("monitoring:maintenance_export")))
        names = [row[1] for row in rows]
        self.assertIn("'=HYPERLINK(\"http://x\")", names)
        self.assertFalse(any(n.startswith("=") for n in names))

    def test_las_fechas_se_exportan_como_fechas(self):
        import datetime
        self.client.force_login(self.viewer)
        rows = self._rows(self.client.get(reverse("monitoring:maintenance_export")))
        self.assertIsInstance(rows[0][8], datetime.datetime)

    def test_el_listado_muestra_el_boton_de_exportar(self):
        self.client.force_login(self.viewer)
        r = self.client.get(reverse("monitoring:maintenance_list"))
        self.assertContains(r, reverse("monitoring:maintenance_export"))
